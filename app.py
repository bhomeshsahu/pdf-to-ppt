from flask import Flask, render_template, request, send_file, redirect, url_for
from pptx import Presentation
from pptx.util import Inches
import fitz  # PyMuPDF
import io
import os

app = Flask(__name__)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/convert', methods=['POST'])
def convert_pdf_to_ppt():
    file = request.files.get('pdf_file')
    if not file:
        return redirect(url_for('index'))

    # Save uploaded PDF temporarily
    pdf_path = "temp.pdf"
    file.save(pdf_path)

    # Create PowerPoint presentation
    prs = Presentation()

    # Open PDF with PyMuPDF
    pdf_document = fitz.open(pdf_path)

    for page_num in range(len(pdf_document)):
        page = pdf_document.load_page(page_num)
        pix = page.get_pixmap(dpi=150)
        img_bytes = pix.tobytes("png")

        # Add slide and insert image
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        image_stream = io.BytesIO(img_bytes)
        slide.shapes.add_picture(image_stream, Inches(0), Inches(0), Inches(10), Inches(7.5))

    pdf_document.close()

    # Save PPTX
    output_path = "converted.pptx"
    prs.save(output_path)

    # Clean up PDF
    os.remove(pdf_path)

    return send_file(output_path, as_attachment=True, download_name="converted.pptx")

if __name__ == '__main__':
    app.run(debug=True)
